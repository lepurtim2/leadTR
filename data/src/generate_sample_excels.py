"""
LeadTR — Sample Excel Generator for B2B Client Outreach
Generates 3 premium, professionally styled Excel sample files (.xlsx)
directly from the DuckDB Parquet lake with clickable WhatsApp links and opportunity tags.

Excludes all government, municipality, public, and institutional entities
to ensure commercial B2B sales quality.
"""

import os
import sys
import duckdb
import openpyxl

sys.stdout.reconfigure(encoding='utf-8')
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
PARQUET_DIR = os.path.join(BASE_DIR, "parquets")
SAMPLES_DIR = os.path.join(os.path.dirname(BASE_DIR), "samples")
os.makedirs(SAMPLES_DIR, exist_ok=True)

# Styling tokens
HEADER_FILL = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")  # Deep Navy Blue
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
ROW_FONT = Font(name="Calibri", size=10, color="1F2937")
LINK_FONT = Font(name="Calibri", size=10, bold=True, color="2563EB", underline="single")
BOLD_FONT = Font(name="Calibri", size=10, bold=True, color="111827")
TAG_HOT_FILL = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")  # Light Red
TAG_HOT_FONT = Font(name="Calibri", size=10, bold=True, color="DC2626")
TAG_ACTIVE_FILL = PatternFill(start_color="D1FAE5", end_color="D1FAE5", fill_type="solid")  # Light Green
TAG_ACTIVE_FONT = Font(name="Calibri", size=10, bold=True, color="059669")

THIN_BORDER = Border(
    left=Side(style='thin', color='E5E7EB'),
    right=Side(style='thin', color='E5E7EB'),
    top=Side(style='thin', color='E5E7EB'),
    bottom=Side(style='thin', color='E5E7EB')
)

# Comprehensive SQL exclusion clause for all public, government & municipal entities
GOV_EXCLUSIONS_SQL = """
    AND NOT (
        -- Genel Devlet & Kamu Kurumları
        canonical_name ILIKE '%devlet%'
        OR canonical_name ILIKE '%belediye%'
        OR canonical_name ILIKE '%belediyesi%'
        OR canonical_name ILIKE '%kaymakam%'
        OR canonical_name ILIKE '%valilik%'
        OR canonical_name ILIKE '%valiliği%'
        OR canonical_name ILIKE '%müdürlüğ%'
        OR canonical_name ILIKE '%bakanlığ%'
        OR canonical_name ILIKE '%muhtar%'
        OR canonical_name ILIKE '%adliye%'
        OR canonical_name ILIKE '%adliyesi%'
        OR canonical_name ILIKE '%emniyet%'
        OR canonical_name ILIKE '%polis%'
        OR canonical_name ILIKE '%jandarma%'
        OR canonical_name ILIKE '%karakol%'
        OR canonical_name ILIKE '%vergi dairesi%'
        OR canonical_name ILIKE '%sosyal güvenlik%'
        OR canonical_name ILIKE '%sgk%'
        OR canonical_name ILIKE '%nüfus%'
        OR canonical_name ILIKE '%milli eğitim%'
        OR canonical_name ILIKE '%cezaevi%'
        OR canonical_name ILIKE '%ceza infaz%'
        OR canonical_name ILIKE '%t.c.%'
        OR canonical_name ILIKE '%kamu%'
        OR canonical_name ILIKE '%resmi%'
        
        -- Kamu Sosyal Tesisleri, Misafirhaneler ve Tesisler
        OR canonical_name ILIKE '%öğretmenevi%'
        OR canonical_name ILIKE '%ogretmenevi%'
        OR canonical_name ILIKE '%polisevi%'
        OR canonical_name ILIKE '%orduevi%'
        OR canonical_name ILIKE '%sosyal tesis%'
        OR canonical_name ILIKE '%sosyal tesisi%'
        OR canonical_name ILIKE '%misafirhane%'
        OR canonical_name ILIKE '%lokal%'
        OR canonical_name ILIKE '%yemekhane%'
        OR canonical_name ILIKE '%kantin%'
        
        -- Sağlık Alanındaki Kamu / Devlet Kurumları
        OR canonical_name ILIKE '%aile sağlığı%'
        OR canonical_name ILIKE '%toplum sağlığı%'
        OR canonical_name ILIKE '%halk sağlığı%'
        OR canonical_name ILIKE '%sağlık ocağ%'
        OR canonical_name ILIKE '%saglik ocag%'
        OR canonical_name ILIKE '%saglık ocag%'
        OR canonical_name ILIKE '%sağlik ocag%'
        OR canonical_name ILIKE '%ilçe sağlık%'
        OR canonical_name ILIKE '%şehir hastanesi%'
        OR canonical_name ILIKE '%üniversite hastanesi%'
        OR canonical_name ILIKE '%eğitim ve araştırma hastanesi%'
        OR canonical_name ILIKE '%eğitim araştırma%'
        OR canonical_name ILIKE '%kamu hastaneler%'
        
        -- Kamu Bankaları, KİT'ler ve Kamu İktisadi Teşekkülleri
        OR canonical_name ILIKE '%ziraat bank%'
        OR canonical_name ILIKE '%vakıfbank%'
        OR canonical_name ILIKE '%halkbank%'
        OR canonical_name ILIKE '%ptt%'
        OR canonical_name ILIKE '%türk telekom%'
        OR canonical_name ILIKE '%iski%'
        OR canonical_name ILIKE '%aski%'
        OR canonical_name ILIKE '%tedaş%'
        OR canonical_name ILIKE '%botaş%'
        OR canonical_name ILIKE '%tmo%'
        OR canonical_name ILIKE '%tarım kredi%'
        OR canonical_name ILIKE '%kızılay%'
        OR canonical_name ILIKE '%yeşilay%'
        
        -- Kategori Kontrolleri
        OR category_name ILIKE '%devlet%'
        OR category_name ILIKE '%kamu%'
        OR category_name ILIKE '%belediye%'
    )
"""

def normalize_province(prov: str, dist: str) -> str:
    """Corrects district-province mappings for pristine Turkish geographic display."""
    dist_clean = (dist or "").strip()
    
    istanbul_districts = {
        'Kadıköy', 'Bakırköy', 'Silivri', 'Pendik', 'Fatih', 'Beşiktaş', 'Şişli', 'Kartal', 
        'Maltepe', 'Ümraniye', 'Üsküdar', 'Beyoğlu', 'Sarıyer', 'Başakşehir', 'Esenyurt', 
        'Sancaktepe', 'Eyüpsultan', 'Zeytinburnu', 'Ataşehir', 'Tuzla', 'Beylikdüzü', 
        'Avcılar', 'Bağcılar', 'Bahçelievler', 'Güngören', 'Bayrampaşa', 'Gaziosmanpaşa', 
        'Sultangazi', 'Arnavutköy', 'Çatalca', 'Büyükçekmece', 'Küçükçekmece', 'Kağıthane'
    }
    ankara_districts = {
        'Çankaya', 'Keçiören', 'Yenimahalle', 'Mamak', 'Etimesgut', 'Sincan', 'Altındağ', 
        'Gölbaşı', 'Polatlı', 'Pursaklar', 'Kahramankazan', 'Beypazarı'
    }
    izmir_districts = {
        'Konak', 'Karşıyaka', 'Bornova', 'Buca', 'Çiğli', 'Bayraklı', 'Balçova', 'Gaziemir', 
        'Urla', 'Çeşme', 'Karabağlar', 'Narlıdere', 'Menemen', 'Torbalı', 'Aliağa', 'Seferihisar'
    }
    bursa_districts = {'Nilüfer', 'Osmangazi', 'Yıldırım', 'İnegöl', 'Gemlik', 'Mudanya', 'Gürsu', 'Kestel'}
    antalya_districts = {'Muratpaşa', 'Kepez', 'Konyaaltı', 'Alanya', 'Manavgat', 'Kemer', 'Serik', 'Kaş'}

    for d in istanbul_districts:
        if d.lower() == dist_clean.lower():
            return 'İstanbul'
    for d in ankara_districts:
        if d.lower() == dist_clean.lower():
            return 'Ankara'
    for d in izmir_districts:
        if d.lower() == dist_clean.lower():
            return 'İzmir'
    for d in bursa_districts:
        if d.lower() == dist_clean.lower():
            return 'Bursa'
    for d in antalya_districts:
        if d.lower() == dist_clean.lower():
            return 'Antalya'

    if not prov or prov in ('Türkiye', 'None', ''):
        return dist_clean if dist_clean else 'Türkiye'
    return prov


def format_workbook(wb, ws, title_text: str):
    """Applies high-end executive styling to the openpyxl worksheet."""
    ws.views.sheetView[0].showGridLines = True
    
    # 1. Title Banner in Row 1
    ws.merge_cells("A1:I1")
    title_cell = ws["A1"]
    title_cell.value = f"📊 LeadTR — {title_text} (Özel Ticari Numune)"
    title_cell.font = Font(name="Calibri", size=14, bold=True, color="1E3A8A")
    title_cell.alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[1].height = 35

    # 2. Subtitle / Disclaimer in Row 2
    ws.merge_cells("A2:I2")
    sub_cell = ws["A2"]
    sub_cell.value = "⚡ Bu liste Türkiye geneli 1.88M doğrulanmış ÖZEL TİCARİ havuzdan seçilmiştir. Kamu/devlet kurumları hariç tutulmuştur. WhatsApp linklerine tıklayarak doğrudan yazabilirsiniz."
    sub_cell.font = Font(name="Calibri", size=9, italic=True, color="6B7280")
    sub_cell.alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[2].height = 20

    # 3. Header Row in Row 4
    ws.row_dimensions[4].height = 26
    for col_idx in range(1, 10):
        cell = ws.cell(row=4, column=col_idx)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = THIN_BORDER

    # 4. Data rows styling
    max_row = ws.max_row
    for r in range(5, max_row + 1):
        ws.row_dimensions[r].height = 22
        for c in range(1, 10):
            cell = ws.cell(row=r, column=c)
            cell.border = THIN_BORDER
            
            # Alignments
            if c in (3, 4, 6, 7, 8, 9):
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

            # Style clickable link
            if c == 7 and cell.value and "wa.me" in str(cell.value):
                cell.font = LINK_FONT

            # Highlight Status Tags
            if c == 8:
                if "Web Sitesi Yok" in str(cell.value):
                    cell.fill = TAG_HOT_FILL
                    cell.font = TAG_HOT_FONT
                else:
                    cell.fill = TAG_ACTIVE_FILL
                    cell.font = TAG_ACTIVE_FONT

    # 5. Auto column widths
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.row in (1, 2):
                continue
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)


def generate_samples():
    con = duckdb.connect()
    con.execute("PRAGMA enable_object_cache=false;")
    lake_pattern = f"{PARQUET_DIR.replace(os.sep, '/')}/sector_*.parquet"

    print("🚀 Generating 3 Tailored Excel Sample Packs for Client Outreach...")
    print("🛡️ Government and public institutions are strictly EXCLUDED.\n")

    # -------------------------------------------------------------
    # PACK 1: Web Tasarım, SEO & Dijital Pazarlama Ajansları İçin
    # Hedef: Telefonu/WhatsApp'ı olan ancak WEB SİTESİ HİÇ OLMAYAN 50 Özel Şirket
    # Hariç: Devlet yerleri, bankalar/ATM, eczaneler
    # -------------------------------------------------------------
    print("[1/3] Generating 'LeadTR_Ornek_Web_Tasarim_ve_SEO_Ajansi_Listesi.xlsx'...")
    rows_web = con.execute(f"""
        SELECT 
            canonical_name,
            category_name,
            province,
            district,
            formatted_address,
            phone,
            lead_score
        FROM '{lake_pattern}'
        WHERE phone IS NOT NULL AND length(trim(phone)) >= 10
          AND (website IS NULL OR length(trim(website)) < 4)
          AND (business_status = 'active' OR business_status IS NULL)
          AND province IN ('İstanbul', 'Ankara', 'İzmir', 'Bursa', 'Antalya', 'Kocaeli', 'Gaziantep', 'Adana', 'Konya', 'Mersin')
          {GOV_EXCLUSIONS_SQL}
          AND NOT (
            -- Bankalar ve Finans Şubeleri
            category_name ILIKE '%banka%' OR canonical_name ILIKE '%bankası%' OR canonical_name ILIKE '%bankamatik%'
            OR canonical_name ILIKE '%atm%' OR canonical_name ILIKE '%garanti%' OR canonical_name ILIKE '%akbank%'
            OR canonical_name ILIKE '%iş bankası%' OR canonical_name ILIKE '%yapı kredi%' OR canonical_name ILIKE '%kuveyt%'
            OR canonical_name ILIKE '%qnb%' OR canonical_name ILIKE '%denizbank%' OR canonical_name ILIKE '%teb%'
            OR canonical_name ILIKE '%finansbank%' OR canonical_name ILIKE '%enpara%'
            
            -- Eczaneler (Web ajansları hedeflemez)
            OR category_name ILIKE '%eczane%' OR canonical_name ILIKE '%eczane%'
            
            -- Hatalı / Çok Kısa İsimler
            OR canonical_name IN ('İşyeri', 'Dükkan', 'Büro', 'Mağaza', 'Ofis')
          )
          AND length(trim(canonical_name)) >= 4
        ORDER BY completeness_score DESC, lead_score DESC
        LIMIT 150
    """).fetchall()

    wb1 = openpyxl.Workbook()
    ws1 = wb1.active
    ws1.title = "Sıcak Web & SEO Leadleri"

    headers = [
        "Özel Firma Unvanı",
        "Sektör / Faaliyet Alanı",
        "İl",
        "İlçe",
        "Açık Adres",
        "Yetkili Telefon",
        "WhatsApp Görüşme Linki",
        "Web Varlık Durumu",
        "Potansiyel Satış Fırsatı"
    ]
    
    for col_idx, h in enumerate(headers, 1):
        ws1.cell(row=4, column=col_idx, value=h)

    seen_names = set()
    row_count = 0
    curr_row = 5

    for r in rows_web:
        name, cat, prov, dist, addr, phone, score = r
        clean_name = name.strip()
        if clean_name.lower() in seen_names:
            continue
        seen_names.add(clean_name.lower())

        clean_digits = "".join(filter(str.isdigit, phone or ""))
        wa_num = clean_digits
        if wa_num.startswith("0"):
            wa_num = "9" + wa_num
        elif not wa_num.startswith("90"):
            wa_num = "90" + wa_num
        
        wa_link = f"https://wa.me/{wa_num}" if len(clean_digits) >= 10 else ""
        norm_prov = normalize_province(prov, dist)

        ws1.cell(row=curr_row, column=1, value=clean_name)
        ws1.cell(row=curr_row, column=2, value=cat or "Özel İşletme")
        ws1.cell(row=curr_row, column=3, value=norm_prov)
        ws1.cell(row=curr_row, column=4, value=dist or "")
        ws1.cell(row=curr_row, column=5, value=addr if addr and len(addr) > 5 else f"{dist}, {norm_prov}")
        ws1.cell(row=curr_row, column=6, value=phone)
        
        c7 = ws1.cell(row=curr_row, column=7)
        if wa_link:
            c7.value = f'=HYPERLINK("{wa_link}", "💬 WhatsApp\'tan Yaz")'
        else:
            c7.value = "-"

        ws1.cell(row=curr_row, column=8, value="❌ Web Sitesi Yok (Sıcak Fırsat)")
        ws1.cell(row=curr_row, column=9, value="🔥 Web Tasarım, SEO & Harita Paketi")

        curr_row += 1
        row_count += 1
        if row_count >= 50:
            break

    format_workbook(wb1, ws1, "Web Tasarım ve Ajanslar İçin Sıcak Leadler")
    file1_path = os.path.join(SAMPLES_DIR, "LeadTR_Ornek_Web_Tasarim_ve_SEO_Ajansi_Listesi.xlsx")
    wb1.save(file1_path)
    print(f" -> Saved: {file1_path} ({row_count} verified private records)")

    # -------------------------------------------------------------
    # PACK 2: Medikal, Diş ve Sağlık Tedarikçileri / B2B Firmalar İçin
    # Hedef: Sadece İNSAN ÖZEL Diş Klinikleri, Muayenehaneler, Özel Poliklinikler
    # Hariç: Tüm devlet hastaneleri, kamu ADSM'leri, ASM, sağlık ocakları, veterinerler
    # -------------------------------------------------------------
    print("\n[2/3] Generating 'LeadTR_Ornek_Dis_Klinikleri_ve_Saglik_Listesi.xlsx'...")
    rows_health = con.execute(f"""
        SELECT 
            canonical_name,
            category_name,
            province,
            district,
            formatted_address,
            phone,
            website,
            lead_score
        FROM '{lake_pattern}'
        WHERE (
            canonical_name ILIKE '%özel%diş%' 
            OR canonical_name ILIKE '%diş hekimi%' 
            OR canonical_name ILIKE '%dt.%' 
            OR canonical_name ILIKE '%dent%' 
            OR canonical_name ILIKE '%dental%' 
            OR canonical_name ILIKE '%ortodonti%' 
            OR canonical_name ILIKE '%implant%'
            OR (canonical_name ILIKE '%poliklinik%' AND (canonical_name ILIKE '%ağız%' OR canonical_name ILIKE '%sağlık%' OR canonical_name ILIKE '%tıp%'))
            OR (category_name ILIKE '%diş%' AND canonical_name NOT ILIKE '%merkez%')
        )
        AND phone IS NOT NULL AND length(trim(phone)) >= 10
        {GOV_EXCLUSIONS_SQL}
        AND NOT (
            (canonical_name ILIKE '%ağız ve diş sağlığı%' OR canonical_name ILIKE '%agiz ve dis sagligi%')
            AND canonical_name NOT ILIKE '%özel%' AND canonical_name NOT ILIKE '%ozel%'
        )
        AND canonical_name NOT ILIKE '%veteriner%'
        AND category_name NOT ILIKE '%veteriner%'
        AND province != 'Türkiye' AND province IS NOT NULL AND length(trim(province)) > 2
        ORDER BY completeness_score DESC, lead_score DESC
        LIMIT 150
    """).fetchall()

    wb2 = openpyxl.Workbook()
    ws2 = wb2.active
    ws2.title = "Özel Klinik & Sağlık Leadleri"

    headers_health = [
        "Özel Sağlık Kuruluşu / Klinik Adı",
        "Uzmanlık Alanı",
        "İl",
        "İlçe",
        "Açık Adres",
        "Randevu / İletişim Hattı",
        "WhatsApp Linki",
        "Web Sitesi",
        "B2B Tedarik Durumu"
    ]
    for col_idx, h in enumerate(headers_health, 1):
        ws2.cell(row=4, column=col_idx, value=h)

    seen_health = set()
    row_count_health = 0
    curr_row = 5

    for r in rows_health:
        name, cat, prov, dist, addr, phone, web, score = r
        clean_name = name.strip()
        if clean_name.lower() in seen_health:
            continue
        seen_health.add(clean_name.lower())

        clean_digits = "".join(filter(str.isdigit, phone or ""))
        wa_num = clean_digits
        if wa_num.startswith("0"):
            wa_num = "9" + wa_num
        elif not wa_num.startswith("90"):
            wa_num = "90" + wa_num
        
        wa_link = f"https://wa.me/{wa_num}" if len(clean_digits) >= 10 else ""
        norm_prov = normalize_province(prov, dist)

        ws2.cell(row=curr_row, column=1, value=clean_name)
        ws2.cell(row=curr_row, column=2, value="Özel Diş Kliniği / Poliklinik")
        ws2.cell(row=curr_row, column=3, value=norm_prov)
        ws2.cell(row=curr_row, column=4, value=dist or "")
        ws2.cell(row=curr_row, column=5, value=addr if addr and len(addr) > 5 else f"{dist}, {norm_prov}")
        ws2.cell(row=curr_row, column=6, value=phone)
        
        c7 = ws2.cell(row=curr_row, column=7)
        if wa_link:
            c7.value = f'=HYPERLINK("{wa_link}", "💬 WhatsApp\'tan Ulaş")'
        else:
            c7.value = "-"

        ws2.cell(row=curr_row, column=8, value=web or "Mevcut Değil")
        ws2.cell(row=curr_row, column=9, value="💉 Tıbbi Sarf Malzeme & Dental Cihaz Satışı")

        curr_row += 1
        row_count_health += 1
        if row_count_health >= 50:
            break

    format_workbook(wb2, ws2, "Özel Diş Klinikleri ve Sağlık Merkezleri B2B Listesi")
    file2_path = os.path.join(SAMPLES_DIR, "LeadTR_Ornek_Dis_Klinikleri_ve_Saglik_Listesi.xlsx")
    wb2.save(file2_path)
    print(f" -> Saved: {file2_path} ({row_count_health} verified private records)")

    # -------------------------------------------------------------
    # PACK 3: Gıda Toptancıları, POS, Adisyon & Endüstriyel Mutfak İçin
    # Hedef: Özel Restoran, Kafe ve Gastronomi İşletmeleri
    # Hariç: Öğretmenevi, polisevi, belediye tesisleri, kamu yemekhaneleri, üniversite kampüsleri
    # -------------------------------------------------------------
    print("\n[3/3] Generating 'LeadTR_Ornek_Restoran_ve_Gastronomi_Listesi.xlsx'...")
    rows_food = con.execute(f"""
        SELECT 
            canonical_name,
            category_name,
            province,
            district,
            formatted_address,
            phone,
            website,
            lead_score
        FROM '{lake_pattern}'
        WHERE (category_name ILIKE '%restoran%' OR category_name ILIKE '%kafe%' OR category_name ILIKE '%restaurant%')
          AND phone IS NOT NULL AND length(trim(phone)) >= 10
          {GOV_EXCLUSIONS_SQL}
          AND province != 'Türkiye' AND province IS NOT NULL AND length(trim(province)) > 2
          AND length(trim(canonical_name)) >= 4
          AND (formatted_address IS NULL OR (
                formatted_address NOT ILIKE '%üniversite%' 
                AND formatted_address NOT ILIKE '%kampüs%' 
                AND formatted_address NOT ILIKE '%yerleşke%' 
                AND formatted_address NOT ILIKE '%kyk%'
          ))
        ORDER BY completeness_score DESC, lead_score DESC
        LIMIT 150
    """).fetchall()

    wb3 = openpyxl.Workbook()
    ws3 = wb3.active
    ws3.title = "Özel Restoran & Kafe Leadleri"

    headers_food = [
        "Özel İşletme Adı",
        "Konsept",
        "İl",
        "İlçe",
        "Açık Adres",
        "Rezervasyon / Sipariş Tel",
        "WhatsApp Sipariş Linki",
        "Dijital Durum",
        "Toptan & POS Fırsatı"
    ]
    for col_idx, h in enumerate(headers_food, 1):
        ws3.cell(row=4, column=col_idx, value=h)

    seen_food = set()
    row_count_food = 0
    curr_row = 5

    for r in rows_food:
        name, cat, prov, dist, addr, phone, web, score = r
        clean_name = name.strip()
        if clean_name.lower() in seen_food:
            continue
        seen_food.add(clean_name.lower())

        clean_digits = "".join(filter(str.isdigit, phone or ""))
        wa_num = clean_digits
        if wa_num.startswith("0"):
            wa_num = "9" + wa_num
        elif not wa_num.startswith("90"):
            wa_num = "90" + wa_num
        
        wa_link = f"https://wa.me/{wa_num}" if len(clean_digits) >= 10 else ""
        norm_prov = normalize_province(prov, dist)

        ws3.cell(row=curr_row, column=1, value=clean_name)
        ws3.cell(row=curr_row, column=2, value=cat or "Restoran & Kafe")
        ws3.cell(row=curr_row, column=3, value=norm_prov)
        ws3.cell(row=curr_row, column=4, value=dist or "")
        ws3.cell(row=curr_row, column=5, value=addr if addr and len(addr) > 5 else f"{dist}, {norm_prov}")
        ws3.cell(row=curr_row, column=6, value=phone)
        
        c7 = ws3.cell(row=curr_row, column=7)
        if wa_link:
            c7.value = f'=HYPERLINK("{wa_link}", "💬 WhatsApp\'tan Ulaş")'
        else:
            c7.value = "-"

        ws3.cell(row=curr_row, column=8, value="✅ Web Sitesi Var" if web else "❌ Web Sitesi Yok")
        ws3.cell(row=curr_row, column=9, value="🍽️ Gıda Tedariği, POS & Adisyon Yazılımı")

        curr_row += 1
        row_count_food += 1
        if row_count_food >= 50:
            break

    format_workbook(wb3, ws3, "Özel Restoran ve Yeme-İçme Sektörü Leadleri")
    file3_path = os.path.join(SAMPLES_DIR, "LeadTR_Ornek_Restoran_ve_Gastronomi_Listesi.xlsx")
    wb3.save(file3_path)
    print(f" -> Saved: {file3_path} ({row_count_food} verified private records)")

    con.close()
    print("\n🎉 ALL 3 SAMPLE EXCEL PACKS RE-GENERATED WITHOUT GOVERNMENT ENTITIES!")


if __name__ == "__main__":
    generate_samples()
